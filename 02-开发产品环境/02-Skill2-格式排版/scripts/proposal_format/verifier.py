from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .parser import ProposalDocument


ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


@dataclass(frozen=True)
class CheckResult:
    name: str
    passed: bool
    expected: Any
    actual: Any

    def as_dict(self) -> dict[str, Any]:
        return {"name": self.name, "passed": self.passed, "expected": self.expected, "actual": self.actual}


def _close(actual: float | None, expected: float, tolerance: float = 0.05) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def _pt(value) -> float | None:
    return None if value is None else round(float(value.pt), 4)


def _mm(value) -> float | None:
    return None if value is None else round(float(value.mm), 4)


def _east_asia(font_owner) -> str | None:
    rpr = font_owner._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    return None if rfonts is None else rfonts.get(qn("w:eastAsia"))


def _paragraph_style_sequence(model: ProposalDocument, profile: dict[str, object]) -> list[str]:
    styles = profile["styles"]
    names = [styles["title"]["word_style"]]
    names.extend(styles[block.type if block.type != "paragraph" else "body"]["word_style"] for block in model.blocks)
    names.extend([styles["signature"]["word_style"], styles["signature"]["word_style"]])
    return [str(name) for name in names]


def verify_docx(path: str | Path, model: ProposalDocument, profile: dict[str, object]) -> list[CheckResult]:
    document = Document(path)
    checks: list[CheckResult] = []

    actual_text = [paragraph.text for paragraph in document.paragraphs]
    expected_text = model.expected_paragraphs()
    checks.append(CheckResult("visible_text", actual_text == expected_text, expected_text, actual_text))

    forbidden = ["【主标题开始】", "【主标题结束】", "【落款】", "## ", "### "]
    combined = "\n".join(actual_text)
    found = [marker for marker in forbidden if marker in combined]
    checks.append(CheckResult("markers_removed", not found, [], found))
    checks.append(CheckResult("no_tables", len(document.tables) == 0, 0, len(document.tables)))
    checks.append(CheckResult("no_images", len(document.inline_shapes) == 0, 0, len(document.inline_shapes)))

    actual_styles = [paragraph.style.name for paragraph in document.paragraphs]
    expected_styles = _paragraph_style_sequence(model, profile)
    checks.append(CheckResult("paragraph_styles", actual_styles == expected_styles, expected_styles, actual_styles))

    section = document.sections[0]
    page = profile["page"]
    page_properties = {
        "page_width_mm": (section.page_width, page["width_mm"]),
        "page_height_mm": (section.page_height, page["height_mm"]),
        "margin_top_mm": (section.top_margin, page["margin_top_mm"]),
        "margin_bottom_mm": (section.bottom_margin, page["margin_bottom_mm"]),
        "margin_left_mm": (section.left_margin, page["margin_left_mm"]),
        "margin_right_mm": (section.right_margin, page["margin_right_mm"]),
        "footer_distance_mm": (section.footer_distance, page["footer_distance_mm"]),
    }
    for name, (actual_value, expected_value) in page_properties.items():
        actual_mm = _mm(actual_value)
        checks.append(CheckResult(name, _close(actual_mm, float(expected_value)), float(expected_value), actual_mm))

    for key, settings in profile["styles"].items():
        style = document.styles[str(settings["word_style"])]
        checks.append(CheckResult(f"{key}.east_asia_font", _east_asia(style) == settings["font_east_asia"], settings["font_east_asia"], _east_asia(style)))
        checks.append(CheckResult(f"{key}.western_font", style.font.name == settings["font_western"], settings["font_western"], style.font.name))
        actual_size = _pt(style.font.size)
        checks.append(CheckResult(f"{key}.size_pt", _close(actual_size, float(settings["size_pt"])), float(settings["size_pt"]), actual_size))
        actual_line = _pt(style.paragraph_format.line_spacing)
        checks.append(CheckResult(f"{key}.line_spacing_pt", _close(actual_line, float(settings["line_spacing_pt"])), float(settings["line_spacing_pt"]), actual_line))
        actual_indent = _pt(style.paragraph_format.first_line_indent) or 0.0
        checks.append(CheckResult(f"{key}.first_line_indent_pt", _close(actual_indent, float(settings["first_line_indent_pt"])), float(settings["first_line_indent_pt"]), actual_indent))
        actual_before = _pt(style.paragraph_format.space_before) or 0.0
        checks.append(CheckResult(f"{key}.space_before_pt", _close(actual_before, float(settings["space_before_pt"])), float(settings["space_before_pt"]), actual_before))
        actual_after = _pt(style.paragraph_format.space_after) or 0.0
        checks.append(CheckResult(f"{key}.space_after_pt", _close(actual_after, float(settings["space_after_pt"])), float(settings["space_after_pt"]), actual_after))
        actual_alignment = style.paragraph_format.alignment
        expected_alignment = ALIGNMENTS[str(settings["alignment"])]
        checks.append(CheckResult(f"{key}.alignment", actual_alignment == expected_alignment, str(settings["alignment"]), str(actual_alignment)))
        actual_bold = bool(style.font.bold)
        checks.append(CheckResult(f"{key}.bold", actual_bold == bool(settings["bold"]), bool(settings["bold"]), actual_bold))

    run_issues: list[str] = []
    western_font = str(profile["fonts"]["western"])
    settings_by_style = {
        str(settings["word_style"]): settings for settings in profile["styles"].values()
    }
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        settings = settings_by_style[paragraph.style.name]
        for run_index, run in enumerate(paragraph.runs, start=1):
            if not run.text:
                continue
            is_numeric = run.text.isdigit()
            expected_east_asia = western_font if is_numeric else str(settings["font_east_asia"])
            expected_western = western_font
            expected_size = float(settings["size_pt"])
            if is_numeric and run.font.name != expected_western:
                run_issues.append(f"p{paragraph_index}r{run_index}:数字西文字体={run.font.name}")
            if _east_asia(run) != expected_east_asia:
                run_issues.append(f"p{paragraph_index}r{run_index}:东亚字体={_east_asia(run)}")
            if not _close(_pt(run.font.size), expected_size):
                run_issues.append(f"p{paragraph_index}r{run_index}:字号={_pt(run.font.size)}")
            if any(char.isdigit() for char in run.text) and not is_numeric:
                run_issues.append(f"p{paragraph_index}r{run_index}:数字未拆分={run.text}")
    checks.append(CheckResult("run_fonts_and_numeric_split", not run_issues, [], run_issues))

    footer = section.footer
    footer_xml = footer._element.xml
    checks.append(CheckResult("page_number_field", "PAGE" in footer_xml, "PAGE field", "present" if "PAGE" in footer_xml else "missing"))
    footer_text = "".join(paragraph.text for paragraph in footer.paragraphs)
    expected_footer = f"{profile['page_number']['left_mark']}{profile['page_number']['start']}{profile['page_number']['right_mark']}"
    checks.append(CheckResult("page_number_text", footer_text == expected_footer, expected_footer, footer_text))
    footer_sizes = sorted({round(run.font.size.pt, 4) for paragraph in footer.paragraphs for run in paragraph.runs if run.font.size})
    expected_size = float(profile["page_number"]["size_pt"])
    checks.append(CheckResult("page_number_size_pt", footer_sizes == [expected_size], [expected_size], footer_sizes))
    footer_alignment = footer.paragraphs[0].alignment
    expected_footer_alignment = ALIGNMENTS[str(profile["page_number"]["alignment"])]
    checks.append(CheckResult("page_number_alignment", footer_alignment == expected_footer_alignment, str(profile["page_number"]["alignment"]), str(footer_alignment)))
    footer_fonts = sorted({run.font.name for paragraph in footer.paragraphs for run in paragraph.runs if run.text})
    expected_footer_font = str(profile["page_number"]["font"])
    checks.append(CheckResult("page_number_font", footer_fonts == [expected_footer_font], [expected_footer_font], footer_fonts))

    expected_portrait = str(page["orientation"]) == "portrait"
    actual_portrait = float(section.page_width) < float(section.page_height)
    checks.append(CheckResult("page_orientation", actual_portrait == expected_portrait, page["orientation"], "portrait" if actual_portrait else "landscape"))

    return checks


def all_passed(checks: list[CheckResult]) -> bool:
    return all(check.passed for check in checks)
